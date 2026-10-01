"""
The Co-Pilot's follow-ups are part of the engagement register.

Each meeting's Co-Pilot questions and follow-ups are checked against the
transcript (core/copilot_followups). Across an engagement, what's still
open from earlier calls is exactly what the register exists to carry.
They are their own section — AI suggestions, not questions anyone
asked — with the register's usual dedupe and provenance.
"""

from __future__ import annotations

from types import SimpleNamespace

from services.engagement_service import EngagementService


def _item(text, kind="clarifying_questions", status="open",
          resolution=None, answer="", seen="2026-10-01T10:00:00"):
    d = {"id": text[:8], "kind": kind, "text": text, "status": status,
         "first_seen": seen, "last_seen": seen, "times_suggested": 1,
         "fresh": False}
    if resolution:
        d["resolution"] = resolution
        d["answer"] = answer
    return d


def _session(sid, board):
    return SimpleNamespace(
        session_id=sid, requirements_struct=[], decisions_struct=[],
        action_items_struct=[], open_questions=[], defects_struct=[],
        copilot_board=board)


def _meta(sid, at):
    return {"session_id": sid, "display_name": f"Call {sid}",
            "started_at": at}


def _reg(*pairs):
    return EngagementService._aggregate(list(pairs))


def test_open_items_from_each_meeting_are_in_the_register():
    reg = _reg(
        (_meta("A", "2026-09-20"), _session("A", [
            _item("Need Salesforce CTI on day one?", resolution="open")])),
        (_meta("B", "2026-10-01"), _session("B", [
            _item("Send the routing workbook to Pat Roe", kind="follow_ups",
                  resolution="partly", answer="Agreed; no owner yet.")])),
    )
    rows = {r["text"]: r for r in reg["copilot_followups"]}
    assert rows["Need Salesforce CTI on day one?"]["kind"] == "question"
    assert rows["Send the routing workbook to Pat Roe"]["kind"] == "follow-up"
    assert all(r["status"] == "open" for r in rows.values())
    assert reg["counts"]["open_copilot_followups"] == 2
    assert all(r["source"] == "co-pilot" for r in rows.values())


def test_answered_in_a_later_call_closes_it_with_the_answer():
    reg = _reg(
        (_meta("A", "2026-09-20"), _session("A", [
            _item("How many agents move in phase one?", resolution="open")])),
        (_meta("B", "2026-10-01"), _session("B", [
            _item("How many agents move in phase one?",
                  resolution="answered", answer="About 400 — Jane Roe.")])),
    )
    [row] = reg["copilot_followups"]
    assert row["status"] == "answered"
    assert row["answer"] == "About 400 — Jane Roe."
    assert [o["session_id"] for o in row["occurrences"]] == ["A", "B"]
    assert reg["counts"]["open_copilot_followups"] == 0


def test_dismissed_items_and_risks_are_left_out():
    reg = _reg((_meta("A", "2026-10-01"), _session("A", [
        _item("Vendor lock-in?", status="dismissed"),
        _item("Porting lead time", kind="risks"),
    ])))
    assert reg["copilot_followups"] == []


def test_handled_during_the_call_is_resolved():
    reg = _reg((_meta("A", "2026-10-01"), _session("A", [
        _item("Ask about SSO", status="done"),
        _item("Send the workbook", kind="follow_ups", status="saved"),
    ])))
    assert {r["status"] for r in reg["copilot_followups"]} == {"done"}


def test_open_items_are_listed_first():
    reg = _reg((_meta("A", "2026-10-01"), _session("A", [
        _item("Answered one", resolution="answered", answer="yes"),
        _item("Still open one", resolution="open"),
    ])))
    assert reg["copilot_followups"][0]["text"] == "Still open one"


def test_a_session_without_a_co_pilot_contributes_nothing():
    s = _session("A", [])
    del s.copilot_board
    reg = _reg((_meta("A", "2026-10-01"), s))
    assert reg["copilot_followups"] == []


def test_the_register_spreadsheet_has_the_sheet_and_the_count(tmp_path):
    import pytest
    pytest.importorskip("openpyxl")   # not in the CI backend venv
    from openpyxl import load_workbook
    from services.engagement_export_service import export_register_workbook
    reg = _reg((_meta("A", "2026-10-01"), _session("A", [
        _item("Need CTI on day one?", resolution="open")])))
    reg.update(client="Acme", project="", generated_at="2026-10-01",
               session_count=1)
    out = export_register_workbook(reg, tmp_path, "register")
    wb = load_workbook(out["path"])
    assert "Co-Pilot Follow-ups" in wb.sheetnames
    rows = list(wb["Co-Pilot Follow-ups"].iter_rows(values_only=True))
    assert any("Need CTI on day one?" in [str(c) for c in r] for r in rows)
    overview = {r[0]: r[1] for r in wb["Overview"].iter_rows(values_only=True)}
    assert overview["Open Co-Pilot follow-ups"] == 1
